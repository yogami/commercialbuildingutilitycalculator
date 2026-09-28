"""Tests for CLI and benchmark reference dataset generator."""

from datetime import date
from decimal import Decimal
from pathlib import Path
import openpyxl
import pytest

from commercial_utility_calculator.application.benchmark import create_spree_campus_benchmark
from commercial_utility_calculator.application.engine import CommercialUtilityEngine
from commercial_utility_calculator.cli import main_cli


def test_spree_campus_benchmark_full_execution(tmp_path: Path) -> None:
    data = create_spree_campus_benchmark()
    assert len(data.spaces) >= 6
    assert len(data.leases) >= 4
    assert len(data.invoices) >= 4

    engine = CommercialUtilityEngine()
    result = engine.calculate(
        spaces=data.spaces,
        meters=data.meters,
        readings=data.readings,
        leases=data.leases,
        invoices=data.invoices,
        billing_start=date(2025, 1, 1),
        billing_end=date(2025, 12, 31),
    )

    # Invariants must hold
    assert result.area_result.total_days == 365
    assert result.area_result.vacancy_sqm_days > Decimal("0.0")  # Vacancy absorbed by landlord

    # Penny balance
    inv_sum = sum((inv.net_amount_eur for inv in data.invoices), Decimal("0.00"))
    alloc_sum = sum((l.net_amount_eur for l in result.cost_lines), Decimal("0.00"))
    assert inv_sum == alloc_sum


def test_cli_benchmark_export(tmp_path: Path) -> None:
    out_excel = tmp_path / "Spree_Campus_Abrechnung_2025.xlsx"
    ret = main_cli(["benchmark", "--output", str(out_excel)])
    assert ret == 0
    assert out_excel.exists()

    wb = openpyxl.load_workbook(str(out_excel), data_only=False)
    assert "Gesamtkosten" in wb.sheetnames
    assert "Flaechen_Uebersicht" in wb.sheetnames


def test_cli_template_generation(tmp_path: Path) -> None:
    tpl_dir = tmp_path / "custom_templates"
    ret = main_cli(["generate-templates", "--output-dir", str(tpl_dir)])
    assert ret == 0
    assert (tpl_dir / "raeume_template.csv").exists()
    assert (tpl_dir / "mietvertraege_template.csv").exists()
    assert (tpl_dir / "rechnungen_template.csv").exists()
