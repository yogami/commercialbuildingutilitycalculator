"""Tests for Phase 5: Invariant verification suite and audit-ready Excel exporter."""

from datetime import date
from decimal import Decimal
from pathlib import Path
import openpyxl
import pytest

from commercial_utility_calculator.domain.entities import (
    CostInvoice,
    Meter,
    MeterReading,
    PhysicalSpace,
    TenantLease,
)
from commercial_utility_calculator.domain.enums import Medium, UsageType
from commercial_utility_calculator.domain.invariants import (
    InvariantValidator,
    PennyBalanceBreachError,
    AreaContinuityBreachError,
    NegativeValueBreachError,
)
from commercial_utility_calculator.application.engine import (
    CommercialUtilityEngine,
    EngineCalculationResult,
)
from commercial_utility_calculator.infrastructure.exporters.excel_exporter import (
    AuditReadyExcelExporter,
)


@pytest.fixture
def benchmark_property() -> tuple[list[PhysicalSpace], list[Meter], list[MeterReading], list[TenantLease], list[CostInvoice]]:
    spaces = [
        PhysicalSpace(
            space_id="SP-A-1",
            building_id="BLDG-A",
            floor=1,
            room_number="101",
            area_sqm=Decimal("200.00"),
            usage_type=UsageType.OFFICE,
            cost_circle_scope=["CAMPUS", "BLDG-A"],
        ),
        PhysicalSpace(
            space_id="SP-A-2",
            building_id="BLDG-A",
            floor=2,
            room_number="201",
            area_sqm=Decimal("300.00"),
            usage_type=UsageType.OFFICE,
            cost_circle_scope=["CAMPUS", "BLDG-A"],
        ),
    ]

    meters = [
        Meter(
            meter_id="M-ROOT-E",
            serial_number="SN-ROOT",
            medium=Medium.ELECTRICITY,
            parent_meter_id=None,
        ),
        Meter(
            meter_id="M-SUB-A1",
            serial_number="SN-SUB1",
            medium=Medium.ELECTRICITY,
            parent_meter_id="M-ROOT-E",
            served_space_ids=["SP-A-1"],
        ),
    ]

    readings = [
        MeterReading(meter_id="M-ROOT-E", reading_date=date(2025, 1, 1), value=Decimal("0.0")),
        MeterReading(meter_id="M-ROOT-E", reading_date=date(2025, 12, 31), value=Decimal("10000.0")),
        MeterReading(meter_id="M-SUB-A1", reading_date=date(2025, 1, 1), value=Decimal("0.0")),
        MeterReading(meter_id="M-SUB-A1", reading_date=date(2025, 12, 31), value=Decimal("6000.0")),
    ]

    leases = [
        TenantLease(
            lease_id="L-1",
            tenant_id="T-1",
            tenant_name="Tech Solutions GmbH",
            space_id="SP-A-1",
            start_date=date(2025, 1, 1),
            end_date=date(2025, 12, 31),
            vat_opt_in=True,
            monthly_prepayment_eur=Decimal("300.00"),
        ),
        TenantLease(
            lease_id="L-2",
            tenant_id="T-2",
            tenant_name="Praxis Dr. Schmidt",
            space_id="SP-A-2",
            start_date=date(2025, 1, 1),
            end_date=date(2025, 6, 30),  # Vacant for second half
            vat_opt_in=False,
            monthly_prepayment_eur=Decimal("200.00"),
        ),
    ]

    invoices = [
        CostInvoice(
            invoice_id="INV-INSURANCE",
            cost_category="Gebaeudeversicherung",
            net_amount_eur=Decimal("2500.00"),
            vat_rate_percent=Decimal("0.0"),  # Insurance is VAT exempt
            gross_amount_eur=Decimal("2500.00"),
            billing_start=date(2025, 1, 1),
            billing_end=date(2025, 12, 31),
            cost_circle_scope="CAMPUS",
        ),
        CostInvoice(
            invoice_id="INV-CLEANING",
            cost_category="Treppenhausreinigung",
            net_amount_eur=Decimal("1200.00"),
            vat_rate_percent=Decimal("19.0"),
            gross_amount_eur=Decimal("1428.00"),
            billing_start=date(2025, 1, 1),
            billing_end=date(2025, 12, 31),
            cost_circle_scope="BLDG-A",
        ),
    ]

    return spaces, meters, readings, leases, invoices


def test_engine_end_to_end_and_invariants(
    benchmark_property: tuple[list[PhysicalSpace], list[Meter], list[MeterReading], list[TenantLease], list[CostInvoice]]
) -> None:
    spaces, meters, readings, leases, invoices = benchmark_property

    engine = CommercialUtilityEngine()
    result = engine.calculate(
        spaces=spaces,
        meters=meters,
        readings=readings,
        leases=leases,
        invoices=invoices,
        billing_start=date(2025, 1, 1),
        billing_end=date(2025, 12, 31),
    )

    # 1. Invariant: Penny balance must hold exact 0.00 EUR difference
    total_invoiced_net = sum((inv.net_amount_eur for inv in invoices), Decimal("0.00"))
    total_allocated_net = sum((line.net_amount_eur for line in result.cost_lines), Decimal("0.00"))
    assert total_allocated_net == total_invoiced_net

    # 2. Invariant: Area continuity
    validator = InvariantValidator()
    validator.assert_area_continuity(result.area_result, spaces, date(2025, 1, 1), date(2025, 12, 31))

    # 3. Invariant: Non-negative values
    validator.assert_non_negative_allocations(result.cost_lines)

    # 4. Check tenant statements generated
    assert "T-1" in result.tenant_statements
    assert "T-2" in result.tenant_statements
    assert "LANDLORD_VACANCY" in result.tenant_statements


def test_excel_exporter_retains_live_formulas(
    benchmark_property: tuple[list[PhysicalSpace], list[Meter], list[MeterReading], list[TenantLease], list[CostInvoice]],
    tmp_path: Path,
) -> None:
    spaces, meters, readings, leases, invoices = benchmark_property

    engine = CommercialUtilityEngine()
    result = engine.calculate(
        spaces=spaces,
        meters=meters,
        readings=readings,
        leases=leases,
        invoices=invoices,
        billing_start=date(2025, 1, 1),
        billing_end=date(2025, 12, 31),
    )

    out_file = tmp_path / "Betriebskostenabrechnung_2025.xlsx"
    exporter = AuditReadyExcelExporter()
    exporter.export_to_excel(result, spaces, str(out_file))

    assert out_file.exists()

    # Load workbook and check formulas in tenant sheets
    wb = openpyxl.load_workbook(str(out_file), data_only=False)
    sheet_names = wb.sheetnames
    assert "Gesamtkosten" in sheet_names
    assert "Flaechen_Uebersicht" in sheet_names
    assert "T-1_Tech Solutions GmbH"[:31] in [s[:31] for s in sheet_names]

    # Verify a cell contains a formula starting with '='
    tenant_ws = wb[sheet_names[2]]
    formula_found = False
    for row in tenant_ws.iter_rows(values_only=False):
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith("="):
                formula_found = True
                break
        if formula_found:
            break

    assert formula_found, "Excel workbook must contain live formulas for audit verification (Prüffähigkeit)"


def test_penny_balance_breach_raises() -> None:
    validator = InvariantValidator()
    invoices = [
        CostInvoice(
            invoice_id="INV-1",
            cost_category="Strom",
            net_amount_eur=Decimal("100.00"),
            vat_rate_percent=Decimal("0.0"),
            gross_amount_eur=Decimal("100.00"),
            billing_start=date(2025, 1, 1),
            billing_end=date(2025, 12, 31),
        )
    ]
    from commercial_utility_calculator.application.services.cost_circle_allocator import CostAllocationLine
    broken_lines = [
        CostAllocationLine(
            invoice_id="INV-1",
            cost_category="Strom",
            cost_circle_scope="CAMPUS",
            tenant_id="T-1",
            tenant_name="Tenant 1",
            space_id="SP-1",
            numerator_sqm_days=Decimal("100"),
            denominator_sqm_days=Decimal("100"),
            cost_ratio=Decimal("1.0"),
            net_amount_eur=Decimal("99.99"),  # 1 cent missing!
            vat_rate_percent=Decimal("0.0"),
            vat_amount_eur=Decimal("0.0"),
            gross_amount_eur=Decimal("99.99"),
            vat_opt_in=True,
            billed_amount_eur=Decimal("99.99"),
        )
    ]
    with pytest.raises(PennyBalanceBreachError, match="Penny balance breached"):
        validator.assert_penny_balance(broken_lines, invoices)


def test_negative_value_breach_raises() -> None:
    validator = InvariantValidator()
    from commercial_utility_calculator.application.services.cost_circle_allocator import CostAllocationLine
    negative_lines = [
        CostAllocationLine(
            invoice_id="INV-1",
            cost_category="Strom",
            cost_circle_scope="CAMPUS",
            tenant_id="T-1",
            tenant_name="Tenant 1",
            space_id="SP-1",
            numerator_sqm_days=Decimal("100"),
            denominator_sqm_days=Decimal("100"),
            cost_ratio=Decimal("1.0"),
            net_amount_eur=Decimal("-10.00"),  # Negative cost!
            vat_rate_percent=Decimal("0.0"),
            vat_amount_eur=Decimal("0.0"),
            gross_amount_eur=Decimal("-10.00"),
            vat_opt_in=True,
            billed_amount_eur=Decimal("-10.00"),
        )
    ]
    with pytest.raises(NegativeValueBreachError, match="Negative"):
        validator.assert_non_negative_allocations(negative_lines)
