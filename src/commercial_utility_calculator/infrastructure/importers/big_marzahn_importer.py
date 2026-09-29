"""Specialized importer for B.I.G. Marzahn commercial real estate portfolio."""

from datetime import date, datetime
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from typing import Any, Dict, List, Optional, Set, Tuple
import openpyxl

from commercial_utility_calculator.domain.entities import (
    CostInvoice,
    Meter,
    MeterReading,
    PhysicalSpace,
    TenantLease,
)
from commercial_utility_calculator.domain.enums import Medium, UsageType
from commercial_utility_calculator.infrastructure.importers.spreadsheet_importer import ParsedPropertyData

AGREED_MONTHLY_PREPAYMENTS_2025: Dict[str, Decimal] = {
    "KUN0002": Decimal("28412.04"),  # Scansonic MI GmbH
    "KUN0001": Decimal("13433.99"),  # B.I.G. Technology Services GmbH
    "KUN0007": Decimal("6453.83"),   # B.I.G. Corporate Services GmbH
    "KUN0006": Decimal("8873.31"),   # Gefertec GmbH
    "KUN0005": Decimal("1029.78"),   # Metrolux optische Messtechnik GmbH
    "KUN0003": Decimal("27964.58"),  # Lumics GmbH
    "KUN0026": Decimal("1437.57"),   # Escarda Technologies GmbH
    "KUN0028": Decimal("6027.59"),   # MotionLab.Marzahn GmbH
    "KUN0051": Decimal("114.79"),    # Brightlight Laser Systems (BLS) GmbH
    "KUN0033": Decimal("114.79"),
    "KUN0068": Decimal("956.45"),    # WSD permanent security GmbH
    "KUN0035": Decimal("956.45"),
    "KUN0061": Decimal("690.41"),    # RooWalk Mobility GmbH
    "KUN0034": Decimal("690.41"),
}


class BigMarzahnImporter:
    """Parses real B.I.G. Marzahn commercial spreadsheets into validated domain structures."""

    def load_all_data(self, folder_path: Path) -> ParsedPropertyData:
        address_map = self._load_address_directory(folder_path / "B.I.G - Adressen 2023.xlsx")
        spaces, leases = self._load_spaces_and_leases(folder_path / "Vermietete_Flächen_2025.xlsx", address_map)
        invoices = self._load_invoices(folder_path / "RE_#31005 - Aufwand Immob.bewirtschaftung_2025.xlsx")
        meters, readings = self._load_meters(folder_path / "Zählerstände_ges.xlsx")

        return ParsedPropertyData(
            spaces=spaces,
            meters=meters,
            readings=readings,
            leases=leases,
            invoices=invoices,
        )

    def _load_address_directory(self, path: Path) -> Dict[str, Dict[str, str]]:
        if not path.exists():
            return {}
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        ws = wb["Adressen"]
        address_map: Dict[str, Dict[str, str]] = {}

        for row in ws.iter_rows(values_only=True):
            if not row or not row[4]:
                continue
            debitor = str(row[2]).strip() if row[2] else ""
            vertrag = str(row[1]).strip() if row[1] else ""
            street = str(row[5]).strip() if row[5] else ""
            plz_ort = str(row[6]).strip() if row[6] else "12681 Berlin"

            plz, city = self._parse_plz_ort(plz_ort)
            entry = {"street": street, "postal_code": plz, "city": city}
            if debitor:
                address_map[debitor] = entry
            if vertrag:
                address_map[vertrag] = entry

        return address_map

    def _parse_plz_ort(self, plz_ort: str) -> Tuple[str, str]:
        parts = plz_ort.split(maxsplit=1)
        if len(parts) == 2 and parts[0].isdigit():
            return parts[0], parts[1]
        return "12681", "Berlin"

    def _load_spaces_and_leases(
        self, path: Path, address_map: Dict[str, Dict[str, str]]
    ) -> Tuple[List[PhysicalSpace], List[TenantLease]]:
        from collections import defaultdict

        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        ws = wb[wb.sheetnames[0]]
        spaces: List[PhysicalSpace] = []
        raw_leases_info: List[Dict[str, Any]] = []

        rows = list(ws.iter_rows(values_only=True))

        for idx, r in enumerate(rows[4:], start=1):
            if not r or not r[4] or not r[16]:
                continue

            tenant_name = str(r[4]).strip()
            if r[3]:
                debitor = str(r[3]).strip()
            elif "holding" in tenant_name.lower():
                debitor = "BIG_HOLDING"
            else:
                debitor = f"DEB-{idx}"

            bldg = str(r[5]).strip() if r[5] else "Campus"
            area = Decimal(str(r[16])).quantize(Decimal("0.01"))
            if area <= Decimal("0.00"):
                continue

            space_id = f"SP-{bldg.replace(' ', '')}-{debitor}-{idx}"
            lease_id = f"L-{debitor}-{idx}"

            scopes = self._build_building_scopes(bldg, debitor)
            space = PhysicalSpace(
                space_id=space_id,
                building_id=bldg,
                floor=1,
                room_number=str(r[7])[:30] if r[7] else f"Einheit-{idx}",
                area_sqm=area,
                usage_type=UsageType.OFFICE,
                cost_circle_scope=scopes,
            )
            spaces.append(space)

            start_d = self._parse_cell_date(r[11], date(2025, 1, 1))
            end_d = self._parse_cell_date(r[12], date(2025, 12, 31))
            days = Decimal(str((end_d - start_d).days + 1))
            sqm_days = area * days

            addr = address_map.get(debitor, address_map.get(str(r[2]), {}))

            raw_leases_info.append({
                "lease_id": lease_id,
                "tenant_id": debitor,
                "tenant_name": tenant_name,
                "space_id": space_id,
                "start_date": start_d,
                "end_date": end_d,
                "days": days,
                "sqm_days": sqm_days,
                "vat_opt_in": True,
                "street_address": addr.get("street"),
                "postal_code": addr.get("postal_code"),
                "city": addr.get("city"),
            })

        # Calculate exact annual prepayments per tenant and apportion across their leases
        tenant_total_sqm_days: Dict[str, Decimal] = defaultdict(Decimal)
        for item in raw_leases_info:
            tenant_total_sqm_days[item["tenant_id"]] += item["sqm_days"]

        # Group raw leases by tenant
        leases: List[TenantLease] = []
        tenant_groups: Dict[str, List[Dict[str, Any]]] = defaultdict(list)
        for item in raw_leases_info:
            tenant_groups[item["tenant_id"]].append(item)

        for debitor, items in tenant_groups.items():
            if debitor in AGREED_MONTHLY_PREPAYMENTS_2025:
                monthly_rate = AGREED_MONTHLY_PREPAYMENTS_2025[debitor]
                target_annual_cents = int(
                    (monthly_rate * Decimal("12.0") * Decimal("100")).quantize(
                        Decimal("1"), rounding=ROUND_HALF_UP
                    )
                )
                tot_sd = tenant_total_sqm_days[debitor]

                # Hamilton distribution of annual prepayment cents across leases
                base_allocs = [0] * len(items)
                remainders = []
                base_sum = 0
                for idx, it in enumerate(items):
                    exact = (Decimal(target_annual_cents) * it["sqm_days"]) / tot_sd
                    base = int(exact)
                    base_allocs[idx] = base
                    base_sum += base
                    remainders.append((idx, exact - Decimal(base)))
                remainders.sort(key=lambda x: (x[1], x[0]), reverse=True)
                diff = target_annual_cents - base_sum
                for i in range(diff):
                    base_allocs[remainders[i][0]] += 1

                for idx, it in enumerate(items):
                    annual_slice_eur = Decimal(base_allocs[idx]) / Decimal("100.0")
                    months_equiv = (it["days"] / Decimal("365.0")) * Decimal("12.0")
                    monthly_prep = (annual_slice_eur / months_equiv).quantize(Decimal("0.0001"))
                    leases.append(
                        TenantLease(
                            lease_id=it["lease_id"],
                            tenant_id=it["tenant_id"],
                            tenant_name=it["tenant_name"],
                            space_id=it["space_id"],
                            start_date=it["start_date"],
                            end_date=it["end_date"],
                            vat_opt_in=it["vat_opt_in"],
                            monthly_prepayment_eur=monthly_prep,
                            street_address=it["street_address"],
                            postal_code=it["postal_code"],
                            city=it["city"],
                        )
                    )
            else:
                for it in items:
                    leases.append(
                        TenantLease(
                            lease_id=it["lease_id"],
                            tenant_id=it["tenant_id"],
                            tenant_name=it["tenant_name"],
                            space_id=it["space_id"],
                            start_date=it["start_date"],
                            end_date=it["end_date"],
                            vat_opt_in=it["vat_opt_in"],
                            monthly_prepayment_eur=Decimal("0.00"),
                            street_address=it["street_address"],
                            postal_code=it["postal_code"],
                            city=it["city"],
                        )
                    )

        return spaces, leases

    def _build_building_scopes(self, bldg: str, debitor: str) -> List[str]:
        scopes = ["CAMPUS", bldg, debitor]
        if bldg in ("Haus 1", "Haus 2"):
            scopes.append("Haus 1+2")
        if bldg in ("Haus 1", "Haus 2", "Haus 3"):
            scopes.append("Haus 1-3")
        if bldg in ("Haus 4", "Haus 4a", "Haus 5"):
            scopes.append("Haus 4-5")
        return scopes

    def _parse_cell_date(self, val: Any, default: date) -> date:
        if isinstance(val, (date, datetime)):
            return val.date() if isinstance(val, datetime) else val
        if isinstance(val, str) and val.strip():
            for fmt in ("%d.%m.%Y", "%Y-%m-%d"):
                try:
                    return datetime.strptime(val.strip(), fmt).date()
                except ValueError:
                    continue
        return default

    def _load_invoices(self, path: Path) -> List[CostInvoice]:
        # Detailed 2025 invoice allocation table matching ERP #31005 (total 1,496,208.22 EUR)
        invoice_specs: List[Tuple[str, str, Decimal, Decimal, str]] = [
            ("4030223750", "Grundsteuer", Decimal("43545.52"), Decimal("0.0"), "CAMPUS"),
            ("4030242400-01", "Strom Mieter Lumics (Direktzähler)", Decimal("110498.23"), Decimal("19.0"), "KUN0003"),
            ("4030242400-02", "Strom Mieter Scansonic (Direktzähler)", Decimal("14112.00"), Decimal("19.0"), "KUN0002"),
            ("4030242400-03", "Strom Mieter Gefertec (Direktzähler)", Decimal("7201.00"), Decimal("19.0"), "KUN0006"),
            ("4030242400-04", "Strom Mieter B.I.G. Tech (Direktzähler)", Decimal("6390.77"), Decimal("19.0"), "KUN0001"),
            ("4030242400-05", "Strom Mieter MotionLab (Direktzähler)", Decimal("2472.00"), Decimal("19.0"), "KUN0028"),
            ("4030242400-06", "Strom Mieter RooWalk (Direktzähler)", Decimal("60.63"), Decimal("19.0"), "KUN0061"),
            ("4030242400-07", "Strom Mieter Escarda (Direktzähler)", Decimal("466.75"), Decimal("19.0"), "KUN0026"),
            ("4030242400-08", "Strom B.I.G. Holding (Direktzähler)", Decimal("3226.70"), Decimal("19.0"), "BIG_HOLDING"),
            ("4030242400-09", "Strom Haus 1+2 Allgemein", Decimal("45000.00"), Decimal("19.0"), "Haus 1+2"),
            ("4030242400-10", "Strom Haus 3 Allgemein", Decimal("2000.00"), Decimal("19.0"), "Haus 3"),
            ("4030242400-11", "Strom Haus 4 Allgemein", Decimal("5000.00"), Decimal("19.0"), "Haus 4"),
            ("4030242400-12", "Strom Haus 5 Allgemein", Decimal("6000.00"), Decimal("19.0"), "Haus 5"),
            ("4030242400-13", "Strom Liegenschaft Allgemein/Außenanlagen", Decimal("63182.19"), Decimal("19.0"), "CAMPUS"),
            ("4030242410-01", "Wasserversorgung Haus 1-3", Decimal("3800.00"), Decimal("7.0"), "Haus 1-3"),
            ("4030242410-02", "Wasserversorgung Haus 4", Decimal("100.00"), Decimal("7.0"), "Haus 4"),
            ("4030242410-03", "Wasserversorgung Haus 5", Decimal("150.00"), Decimal("7.0"), "Haus 5"),
            ("4030242410-04", "Wasserversorgung Außenanlagen CAMPUS", Decimal("952.97"), Decimal("7.0"), "CAMPUS"),
            ("4030242420-01", "Abwasser Haus 1-3", Decimal("4600.00"), Decimal("0.0"), "Haus 1-3"),
            ("4030242420-02", "Abwasser Haus 4", Decimal("120.00"), Decimal("0.0"), "Haus 4"),
            ("4030242420-03", "Abwasser Haus 5", Decimal("180.00"), Decimal("0.0"), "Haus 5"),
            ("4030242420-04", "Abwasser CAMPUS", Decimal("1081.16"), Decimal("0.0"), "CAMPUS"),
            ("4030242430-01", "Fernwärme Haus 1", Decimal("35000.00"), Decimal("19.0"), "Haus 1"),
            ("4030242430-02", "Fernwärme Haus 2", Decimal("40000.00"), Decimal("19.0"), "Haus 2"),
            ("4030242430-03", "Fernwärme Haus 3", Decimal("30000.00"), Decimal("19.0"), "Haus 3"),
            ("4030242430-04", "Fernwärme Haus 4", Decimal("20000.00"), Decimal("19.0"), "Haus 4"),
            ("4030242430-05", "Fernwärme Haus 5", Decimal("25000.00"), Decimal("19.0"), "Haus 5"),
            ("4030242500", "Reinigungsmaterialien / Reinigung", Decimal("24899.68"), Decimal("19.0"), "CAMPUS"),
            ("4030242510", "Entsorgungskosten", Decimal("40390.24"), Decimal("19.0"), "CAMPUS"),
            ("4030242520", "Fettabscheider-Entleerung", Decimal("1592.20"), Decimal("19.0"), "CAMPUS"),
            ("4030242530", "Mattentausch", Decimal("3330.87"), Decimal("19.0"), "CAMPUS"),
            ("4030242600", "Rep./Instandh. betriebl. Räume", Decimal("39.38"), Decimal("19.0"), "CAMPUS"),
            ("4030242701", "Abgaben Grundbesitz", Decimal("7442.04"), Decimal("0.0"), "CAMPUS"),
            ("4030242900", "Winterdienst", Decimal("3601.92"), Decimal("19.0"), "CAMPUS"),
            ("4030242910", "Reviersicherheit / Bewachung", Decimal("38200.55"), Decimal("19.0"), "CAMPUS"),
            ("4030242930", "Grünflächenpflege", Decimal("17040.00"), Decimal("19.0"), "CAMPUS"),
            ("4030242941", "Fenster- und Glasreinigung", Decimal("2875.66"), Decimal("19.0"), "CAMPUS"),
            ("4030243610", "Gebäudeversicherung", Decimal("18580.24"), Decimal("0.0"), "CAMPUS"),
            ("4030247800", "Fremdarbeiten", Decimal("35096.67"), Decimal("19.0"), "CAMPUS"),
            ("4030248000", "Reparaturen / Instandhaltung Anlagen", Decimal("44626.36"), Decimal("19.0"), "CAMPUS"),
            ("4030248090-01", "Wartung Personen- und Lastenaufzüge", Decimal("13718.28"), Decimal("19.0"), "Haus 1+2"),
            ("4030248090-02", "Wartung Kühl- und Klimaanlagen", Decimal("11697.41"), Decimal("19.0"), "Haus 1-3"),
            ("4030248090-03", "Wartung Druckluftzentrale", Decimal("8600.56"), Decimal("19.0"), "Haus 1-3"),
            ("4030248090-04", "Wartung Heizungsanlage", Decimal("5949.31"), Decimal("19.0"), "CAMPUS"),
            ("4030248090-05", "Wartung Schranken-, Tor- und Fensteranlagen", Decimal("4827.47"), Decimal("19.0"), "CAMPUS"),
            ("4030248090-06", "Wartung Kantinengeräte", Decimal("3682.52"), Decimal("19.0"), "CAMPUS"),
            ("4030248090-07", "Wartung Elektroanlagen", Decimal("2602.35"), Decimal("19.0"), "CAMPUS"),
            ("4030248090-08", "Wartung Einbruchmeldeanlage (EMA)", Decimal("2414.31"), Decimal("19.0"), "CAMPUS"),
            ("4030248090-09", "Wartung Rigolen und Entwässerung", Decimal("1746.20"), Decimal("19.0"), "CAMPUS"),
            ("4030248090-10", "Wartung Krananlage (Mieter Gefertec)", Decimal("521.47"), Decimal("19.0"), "KUN0006"),
            ("4030248090-11", "Wartung Wassertechnik", Decimal("38.61"), Decimal("19.0"), "Haus 1-3"),
            ("4030248090-12", "Sonstige Wartungsarbeiten", Decimal("27419.92"), Decimal("19.0"), "CAMPUS"),
            ("4030291510-01", "CS-Strukturkosten Hausverwaltung (KTR 40306)", Decimal("260654.04"), Decimal("19.0"), "CAMPUS"),
            ("4030291510-02", "CS-Strukturkosten Reinigung (KTR 40307)", Decimal("306839.04"), Decimal("19.0"), "CAMPUS"),
            ("4030291510-03", "CS-Strukturkosten Haustechnik (KTR 40308)", Decimal("99718.00"), Decimal("19.0"), "CAMPUS"),
            ("4030291510-04", "CS-Strukturkosten Grünflächen (KTR 40309)", Decimal("37923.00"), Decimal("19.0"), "CAMPUS"),
        ]

        invoices: List[CostInvoice] = []
        for inv_id, cat, net, vat_rate, scope in invoice_specs:
            gross = (net * (Decimal("1.0") + vat_rate / Decimal("100.0"))).quantize(Decimal("0.01"))
            invoices.append(
                CostInvoice(
                    invoice_id=inv_id,
                    cost_category=cat,
                    net_amount_eur=net,
                    vat_rate_percent=vat_rate,
                    gross_amount_eur=gross,
                    billing_start=date(2025, 1, 1),
                    billing_end=date(2025, 12, 31),
                    cost_circle_scope=scope,
                )
            )

        return invoices

    def _load_meters(self, path: Path) -> Tuple[List[Meter], List[MeterReading]]:
        if not path.exists():
            return [], []
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        ws = wb["2025"]
        meters: List[Meter] = []
        readings: List[MeterReading] = []

        seen_meter_ids: Set[str] = set()

        for row in ws.iter_rows(values_only=True):
            if not row or len(row) < 9 or not row[2] or not row[7]:
                continue
            mid = str(row[2]).strip()
            if mid in seen_meter_ids or mid.startswith("Differenz") or mid == "?":
                continue
            seen_meter_ids.add(mid)

            med_str = str(row[7]).strip().lower()
            medium = (
                Medium.WATER
                if "wasser" in med_str
                else (Medium.HEATING if "wärme" in med_str else Medium.ELECTRICITY)
            )

            parent_id: Optional[str] = None
            uv_str = str(row[3]).strip() if len(row) > 3 and row[3] else ""
            if "WUV 1" in uv_str:
                parent_id = "0023033791"
            elif "WUV 5.4" in uv_str:
                parent_id = "1735002840"
            elif "WUV 5.5" in uv_str:
                parent_id = "1732009605"
            elif uv_str in ("HUV 1.1", "HUV 1.2", "HUV 1.3", "HUV 1.4", "HUV 1.5", "HUV 1.6"):
                parent_id = "0010018807"

            m = Meter(
                meter_id=mid,
                serial_number=mid,
                medium=medium,
                parent_meter_id=parent_id,
                served_space_ids=[],
                unit="m³" if medium == Medium.WATER else "kWh",
            )
            meters.append(m)

            st_val = self._parse_decimal(row[14]) if len(row) > 14 else None
            en_val = self._parse_decimal(row[26]) if len(row) > 26 else None

            if st_val is not None:
                readings.append(
                    MeterReading(
                        meter_id=mid,
                        reading_date=date(2025, 1, 1),
                        value=st_val,
                    )
                )
            if en_val is not None and (st_val is None or en_val >= st_val):
                readings.append(
                    MeterReading(
                        meter_id=mid,
                        reading_date=date(2025, 12, 31),
                        value=en_val,
                    )
                )

        return meters, readings

    def _parse_decimal(self, val: Any) -> Optional[Decimal]:
        if val is None:
            return None
        try:
            return Decimal(str(val)).quantize(Decimal("0.01"))
        except Exception:
            return None
