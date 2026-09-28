"""Ingestion parser supporting German commercial Excel and CSV formats."""

import csv
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any, Dict, List, Optional
import openpyxl

from commercial_utility_calculator.domain.entities import (
    CostInvoice,
    Meter,
    MeterReading,
    PhysicalSpace,
    TenantLease,
)
from commercial_utility_calculator.domain.enums import Medium, UsageType


@dataclass
class ParsedPropertyData:
    spaces: List[PhysicalSpace] = field(default_factory=list)
    meters: List[Meter] = field(default_factory=list)
    readings: List[MeterReading] = field(default_factory=list)
    leases: List[TenantLease] = field(default_factory=list)
    invoices: List[CostInvoice] = field(default_factory=list)


class SpreadsheetImporter:
    """Parses German legacy spreadsheets or CSV exports into validated domain models."""

    def import_from_csvs(
        self,
        spaces_path: Optional[str] = None,
        leases_path: Optional[str] = None,
        invoices_path: Optional[str] = None,
        meters_path: Optional[str] = None,
        readings_path: Optional[str] = None,
    ) -> ParsedPropertyData:
        data = ParsedPropertyData()
        if spaces_path and Path(spaces_path).exists():
            data.spaces = self._parse_spaces_csv(spaces_path)
        if leases_path and Path(leases_path).exists():
            data.leases = self._parse_leases_csv(leases_path)
        if invoices_path and Path(invoices_path).exists():
            data.invoices = self._parse_invoices_csv(invoices_path)
        if meters_path and Path(meters_path).exists():
            data.meters = self._parse_meters_csv(meters_path)
        if readings_path and Path(readings_path).exists():
            data.readings = self._parse_readings_csv(readings_path)
        return data

    def import_from_excel(self, excel_path: str) -> ParsedPropertyData:
        wb = openpyxl.load_workbook(excel_path, data_only=True)
        data = ParsedPropertyData()

        spaces_sheet = self._find_sheet(wb, ["raeume", "spaces", "flaechen"])
        if spaces_sheet:
            data.spaces = self._parse_spaces_rows(self._sheet_to_dicts(spaces_sheet))

        leases_sheet = self._find_sheet(wb, ["mietvertraege", "leases", "mieter"])
        if leases_sheet:
            data.leases = self._parse_leases_rows(self._sheet_to_dicts(leases_sheet))

        invoices_sheet = self._find_sheet(wb, ["rechnungen", "invoices", "kosten"])
        if invoices_sheet:
            data.invoices = self._parse_invoices_rows(self._sheet_to_dicts(invoices_sheet))

        meters_sheet = self._find_sheet(wb, ["zaehler", "meters"])
        if meters_sheet:
            data.meters = self._parse_meters_rows(self._sheet_to_dicts(meters_sheet))

        readings_sheet = self._find_sheet(wb, ["zaehlerstaende", "readings", "ablesungen"])
        if readings_sheet:
            data.readings = self._parse_readings_rows(self._sheet_to_dicts(readings_sheet))

        return data

    def _parse_spaces_csv(self, path: str) -> List[PhysicalSpace]:
        rows = self._read_csv_dicts(path)
        return self._parse_spaces_rows(rows)

    def _parse_spaces_rows(self, rows: List[Dict[str, Any]]) -> List[PhysicalSpace]:
        spaces: List[PhysicalSpace] = []
        for r in rows:
            sid = self._get_str(r, ["space_id", "raum_id", "einheit_id"])
            bid = self._get_str(r, ["building_id", "gebaeude_id", "haus"])
            fl = int(self._get_num(r, ["floor", "etage", "stockwerk"], default="0"))
            rm = self._get_str(r, ["room_number", "raumnummer", "nr"])
            sqm = self._parse_decimal(self._get_val(r, ["area_sqm", "flaeche_qm", "flaeche"]))
            usage = self._parse_usage_type(self._get_str(r, ["usage_type", "nutzungsart"]))
            scopes = self._parse_scopes(self._get_str(r, ["cost_circle_scope", "kostenkreis", "scope"]))

            spaces.append(PhysicalSpace(
                space_id=sid,
                building_id=bid,
                floor=fl,
                room_number=rm,
                area_sqm=sqm,
                usage_type=usage,
                cost_circle_scope=scopes,
            ))
        return spaces

    def _parse_leases_csv(self, path: str) -> List[TenantLease]:
        rows = self._read_csv_dicts(path)
        return self._parse_leases_rows(rows)

    def _parse_leases_rows(self, rows: List[Dict[str, Any]]) -> List[TenantLease]:
        leases: List[TenantLease] = []
        for r in rows:
            lid = self._get_str(r, ["lease_id", "vertrag_id"])
            tid = self._get_str(r, ["tenant_id", "mieter_id"])
            tname = self._get_str(r, ["tenant_name", "mieter_name", "mieter"])
            sid = self._get_str(r, ["space_id", "raum_id", "einheit_id"])
            start_d = self._parse_date(self._get_val(r, ["start_date", "mietbeginn", "von"]))
            end_d = self._parse_date(self._get_val(r, ["end_date", "mietende", "bis"]))
            vat_opt = self._parse_bool(self._get_val(r, ["vat_opt_in", "ust_option", "vorsteuerabzug"]), default=True)
            prep = self._parse_decimal(self._get_val(r, ["monthly_prepayment_eur", "vorauszahlung_monatlich", "vorauszahlung"], default="0.00"))

            leases.append(TenantLease(
                lease_id=lid,
                tenant_id=tid,
                tenant_name=tname,
                space_id=sid,
                start_date=start_d,
                end_date=end_d,
                vat_opt_in=vat_opt,
                monthly_prepayment_eur=prep,
            ))
        return leases

    def _parse_invoices_csv(self, path: str) -> List[CostInvoice]:
        rows = self._read_csv_dicts(path)
        return self._parse_invoices_rows(rows)

    def _parse_invoices_rows(self, rows: List[Dict[str, Any]]) -> List[CostInvoice]:
        invoices: List[CostInvoice] = []
        for r in rows:
            iid = self._get_str(r, ["invoice_id", "rechnungs_id", "nr"])
            cat = self._get_str(r, ["cost_category", "kostenart", "kategorie"])
            aid = self._get_str(r, ["asset_id", "anlagen_id"], default=None)
            net = self._parse_decimal(self._get_val(r, ["net_amount_eur", "betrag_netto", "netto"]))
            vat_rate = self._parse_decimal(self._get_val(r, ["vat_rate_percent", "ust_satz", "mwst"], default="19.0"))
            gross = self._parse_decimal(self._get_val(r, ["gross_amount_eur", "betrag_brutto", "brutto"], default=None) or str(net * (Decimal("1.0") + vat_rate / Decimal("100.0"))))
            b_start = self._parse_date(self._get_val(r, ["billing_start", "zeitraum_start", "von"]))
            b_end = self._parse_date(self._get_val(r, ["billing_end", "zeitraum_ende", "bis"]))
            scope = self._get_str(r, ["cost_circle_scope", "kostenkreis", "scope"], default="CAMPUS")

            invoices.append(CostInvoice(
                invoice_id=iid,
                cost_category=cat,
                asset_id=aid if aid else None,
                net_amount_eur=net,
                vat_rate_percent=vat_rate,
                gross_amount_eur=gross,
                billing_start=b_start,
                billing_end=b_end,
                cost_circle_scope=scope,
            ))
        return invoices

    def _parse_meters_csv(self, path: str) -> List[Meter]:
        rows = self._read_csv_dicts(path)
        return self._parse_meters_rows(rows)

    def _parse_meters_rows(self, rows: List[Dict[str, Any]]) -> List[Meter]:
        meters: List[Meter] = []
        for r in rows:
            mid = self._get_str(r, ["meter_id", "zaehler_id"])
            sn = self._get_str(r, ["serial_number", "seriennummer"])
            medium = self._parse_medium(self._get_str(r, ["medium"]))
            pmid = self._get_str(r, ["parent_meter_id", "uebergeordneter_zaehler_id"], default=None)
            spaces_raw = self._get_str(r, ["served_space_ids", "zugeordnete_raeume"], default="")
            spaces = [s.strip() for s in spaces_raw.split(";") if s.strip()] if spaces_raw else []
            mult = self._parse_decimal(self._get_val(r, ["multiplier", "wandlerfaktor"], default="1.0"))
            unit = self._get_str(r, ["unit", "einheit"], default="kWh")

            meters.append(Meter(
                meter_id=mid,
                serial_number=sn,
                medium=medium,
                parent_meter_id=pmid if pmid else None,
                served_space_ids=spaces,
                multiplier=mult,
                unit=unit,
            ))
        return meters

    def _parse_readings_csv(self, path: str) -> List[MeterReading]:
        rows = self._read_csv_dicts(path)
        return self._parse_readings_rows(rows)

    def _parse_readings_rows(self, rows: List[Dict[str, Any]]) -> List[MeterReading]:
        readings: List[MeterReading] = []
        for r in rows:
            mid = self._get_str(r, ["meter_id", "zaehler_id"])
            rd = self._parse_date(self._get_val(r, ["reading_date", "ablesedatum", "datum"]))
            val = self._parse_decimal(self._get_val(r, ["value", "stand", "zaehlerstand"]))
            readings.append(MeterReading(meter_id=mid, reading_date=rd, value=val))
        return readings

    def _parse_decimal(self, raw_val: Any) -> Decimal:
        if raw_val is None:
            return Decimal("0.00")
        s = str(raw_val).strip().replace("€", "").replace("m²", "").replace("kWh", "").replace(" ", "")
        s = s.replace(".", "").replace(",", ".") if "," in s and s.count(".") <= 1 else s.replace(",", ".")
        return Decimal(s)

    def _parse_date(self, raw_val: Any) -> date:
        if isinstance(raw_val, (date, datetime)):
            return raw_val.date() if isinstance(raw_val, datetime) else raw_val
        s = str(raw_val).strip()
        for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y"):
            try:
                return datetime.strptime(s, fmt).date()
            except ValueError:
                continue
        raise ValueError(f"Unable to parse date string: {raw_val}")

    def _parse_bool(self, raw_val: Any, default: bool = True) -> bool:
        if raw_val is None:
            return default
        s = str(raw_val).strip().lower()
        if s in ("ja", "yes", "true", "1", "optiert"):
            return True
        if s in ("nein", "no", "false", "0", "befreit"):
            return False
        return default

    def _parse_usage_type(self, raw_val: str) -> UsageType:
        s = (raw_val or "").strip().lower()
        if "praxis" in s or "arzt" in s or "med" in s:
            return UsageType.MEDICAL
        if "lager" in s or "storage" in s:
            return UsageType.STORAGE
        if "kantine" in s or "canteen" in s or "kueche" in s:
            return UsageType.CANTEEN
        if "labor" in s or "gewerbe" in s or "industrial" in s:
            return UsageType.LIGHT_INDUSTRIAL
        if "allgemein" in s or "common" in s:
            return UsageType.COMMON_AREA
        return UsageType.OFFICE

    def _parse_medium(self, raw_val: str) -> Medium:
        s = (raw_val or "").strip().lower()
        if "strom" in s or "elec" in s:
            return Medium.ELECTRICITY
        if "wasser" in s or "water" in s:
            return Medium.WATER
        if "waerme" in s or "wärme" in s or "heat" in s or "fernwaerme" in s:
            return Medium.HEATING
        if "kaelte" in s or "kälte" in s or "cool" in s:
            return Medium.COOLING
        return Medium.ELECTRICITY

    def _parse_scopes(self, raw_val: str) -> List[str]:
        if not raw_val:
            return ["CAMPUS"]
        scopes = [part.strip() for part in raw_val.replace(",", ";").split(";") if part.strip()]
        return scopes or ["CAMPUS"]

    def _read_csv_dicts(self, path: str) -> List[Dict[str, Any]]:
        with open(path, mode="r", encoding="utf-8-sig") as f:
            sample = f.read(2048)
            f.seek(0)
            delimiter = ";" if sample.count(";") > sample.count(",") else ","
            reader = csv.DictReader(f, delimiter=delimiter)
            return list(reader)

    def _sheet_to_dicts(self, sheet: openpyxl.worksheet.worksheet.Worksheet) -> List[Dict[str, Any]]:
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [str(h).strip() if h is not None else f"col_{idx}" for idx, h in enumerate(rows[0])]
        result: List[Dict[str, Any]] = []
        for r in rows[1:]:
            if not any(r):
                continue
            row_dict = {headers[i]: r[i] for i in range(min(len(headers), len(r)))}
            result.append(row_dict)
        return result

    def _find_sheet(self, wb: openpyxl.Workbook, keywords: List[str]) -> Optional[openpyxl.worksheet.worksheet.Worksheet]:
        for name in wb.sheetnames:
            clean = name.lower().replace(" ", "").replace("_", "")
            if any(k in clean for k in keywords):
                return wb[name]
        return None

    def _get_val(self, row: Dict[str, Any], candidate_keys: List[str], default: Any = None) -> Any:
        norm = {str(k).lower().replace(" ", "").replace("_", ""): v for k, v in row.items() if k is not None}
        for ck in candidate_keys:
            clean_ck = ck.lower().replace(" ", "").replace("_", "")
            if clean_ck in norm and norm[clean_ck] is not None:
                return norm[clean_ck]
        return default

    def _get_str(self, row: Dict[str, Any], candidate_keys: List[str], default: Any = "") -> str:
        val = self._get_val(row, candidate_keys, default)
        return str(val).strip() if val is not None else ""

    def _get_num(self, row: Dict[str, Any], candidate_keys: List[str], default: Any = "0") -> str:
        val = self._get_val(row, candidate_keys, default)
        return str(val).strip() if val is not None else "0"
